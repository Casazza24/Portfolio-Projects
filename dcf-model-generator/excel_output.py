import io

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from datetime import datetime

SCALE = 1_000_000

# -- Colour palette --
DARK_HEADER = PatternFill("solid", fgColor="1F4E79")
SECTION_FILL = PatternFill("solid", fgColor="D6E4F0")
LIGHT_BLUE = PatternFill("solid", fgColor="DAEEF3")
SUBTOTAL_FILL = PatternFill("solid", fgColor="F2F2F2")
GREEN_FILL = PatternFill("solid", fgColor="C6EFCE")
RED_FILL = PatternFill("solid", fgColor="FFC7CE")
WHITE_FILL = PatternFill("solid", fgColor="FFFFFF")

# -- Fonts --
TITLE_FONT = Font(name="Calibri", size=14, bold=True, color="1F4E79")
SUBTITLE_FONT = Font(name="Calibri", size=10, italic=True, color="808080")
SECTION_FONT = Font(name="Calibri", size=11, bold=True, color="1F4E79")
HEADER_FONT = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
LABEL_FONT = Font(name="Calibri", size=10)
BOLD_FONT = Font(name="Calibri", size=10, bold=True)
DATA_FONT = Font(name="Calibri", size=10)

# -- Borders --
THIN = Side(style="thin", color="B0B0B0")
MEDIUM = Side(style="medium", color="1F4E79")
HAIR = Side(style="hair", color="D0D0D0")

CELL_BORDER = Border(bottom=HAIR)
HEADER_BORDER = Border(bottom=MEDIUM)
SUBTOTAL_BORDER = Border(top=THIN, bottom=Side(style="double", color="1F4E79"))
SECTION_BOTTOM = Border(bottom=MEDIUM)
BOX_TOP = Border(top=MEDIUM, left=MEDIUM, right=MEDIUM)
BOX_MID = Border(left=MEDIUM, right=MEDIUM)
BOX_BOT = Border(bottom=MEDIUM, left=MEDIUM, right=MEDIUM)
BOX_SINGLE = Border(top=MEDIUM, bottom=MEDIUM, left=MEDIUM, right=MEDIUM)

RIGHT = Alignment(horizontal="right", vertical="center")
LEFT = Alignment(horizontal="left", vertical="center")
CENTER = Alignment(horizontal="center", vertical="center")

PCT_ROWS = {"Gross Margin", "EBIT Margin", "Revenue Growth"}
BOLD_ROWS = {"Total Revenue", "Gross Profit", "EBIT", "EBITDA",
             "Net Income", "NOPAT", "UFCF", "Free Cash Flow"}


def _millions(val):
    """Scale a raw value to millions."""
    if pd.isna(val):
        return None
    return val / SCALE


def _set_col_widths(ws, widths):
    """Set column widths from a dict of {col_letter: width}."""
    for letter, w in widths.items():
        ws.column_dimensions[letter].width = w


def _auto_fit(ws, min_width=10, max_width=50):
    for col_cells in ws.columns:
        col_letter = col_cells[0].column_letter
        length = max(
            (len(str(c.value)) for c in col_cells if c.value is not None),
            default=min_width,
        )
        ws.column_dimensions[col_letter].width = min(max(length + 3, min_width), max_width)


def _write_summary(wb, ticker, company_name, current_price, implied_price, wacc, assumptions):
    ws = wb.active
    ws.title = "Summary"
    ws.sheet_properties.tabColor = "1F4E79"

    # Title
    ws.merge_cells("A1:D1")
    title_cell = ws["A1"]
    title_cell.value = f"{company_name} ({ticker})"
    title_cell.font = Font(name="Calibri", size=16, bold=True, color="1F4E79")
    title_cell.alignment = LEFT

    ws.merge_cells("A2:D2")
    ws["A2"].value = f"Discounted Cash Flow Analysis  |  {datetime.now().strftime('%B %d, %Y')}"
    ws["A2"].font = SUBTITLE_FONT
    ws["A2"].alignment = LEFT

    # --- Valuation Summary box ---
    row = 4
    ws.merge_cells(f"A{row}:B{row}")
    ws[f"A{row}"].value = "Valuation Summary"
    ws[f"A{row}"].font = SECTION_FONT
    ws[f"A{row}"].border = SECTION_BOTTOM
    ws[f"B{row}"].border = SECTION_BOTTOM
    row += 1

    upside = (implied_price - current_price) / current_price

    val_items = [
        ("Current Market Price", current_price, "$#,##0.00"),
        ("Implied Share Price", implied_price, "$#,##0.00"),
        ("Upside / Downside", upside, "0.0%"),
    ]
    for label, value, fmt in val_items:
        ws.cell(row=row, column=1, value=label).font = LABEL_FONT
        ws.cell(row=row, column=1).alignment = LEFT
        c = ws.cell(row=row, column=2, value=value)
        c.number_format = fmt
        c.font = BOLD_FONT
        c.alignment = RIGHT
        if label == "Upside / Downside":
            c.font = Font(name="Calibri", size=10, bold=True,
                          color="006100" if upside >= 0 else "9C0006")
        row += 1

    # --- Key Assumptions box ---
    row += 1
    ws.merge_cells(f"A{row}:B{row}")
    ws[f"A{row}"].value = "Key Assumptions"
    ws[f"A{row}"].font = SECTION_FONT
    ws[f"A{row}"].border = SECTION_BOTTOM
    ws[f"B{row}"].border = SECTION_BOTTOM
    row += 1

    assumption_items = [
        ("WACC", wacc, "0.00%"),
        ("Terminal Growth Rate", assumptions["terminal_growth_rate"], "0.00%"),
        ("Tax Rate", assumptions["tax_rate"], "0.0%"),
        ("Gross Margin (Start)", assumptions["gross_margins"][0], "0.0%"),
        ("Gross Margin (Target)", assumptions["gross_margins"][-1], "0.0%"),
        ("SG&A % Revenue", assumptions["sga_pct_revenue"], "0.0%"),
        ("D&A % Revenue", assumptions["da_pct_revenue"], "0.0%"),
        ("CapEx % Revenue", assumptions["capex_pct_revenue"], "0.0%"),
        ("NWC % Revenue", assumptions["nwc_pct_revenue"], "0.0%"),
    ]
    for label, value, fmt in assumption_items:
        ws.cell(row=row, column=1, value=label).font = LABEL_FONT
        ws.cell(row=row, column=1).alignment = LEFT
        c = ws.cell(row=row, column=2, value=value)
        c.number_format = fmt
        c.font = DATA_FONT
        c.alignment = RIGHT
        row += 1

    # --- Growth schedule ---
    row += 1
    ws.merge_cells(f"A{row}:B{row}")
    ws[f"A{row}"].value = "Revenue Growth Schedule"
    ws[f"A{row}"].font = SECTION_FONT
    ws[f"A{row}"].border = SECTION_BOTTOM
    ws[f"B{row}"].border = SECTION_BOTTOM
    row += 1
    for i, g in enumerate(assumptions["revenue_growth_rates"]):
        ws.cell(row=row, column=1, value=f"Year {i+1}").font = LABEL_FONT
        c = ws.cell(row=row, column=2, value=g)
        c.number_format = "0.0%"
        c.font = DATA_FONT
        c.alignment = RIGHT
        row += 1

    _set_col_widths(ws, {"A": 24, "B": 16, "C": 4, "D": 4})


def _write_financials(wb, full_df):
    ws = wb.create_sheet("Financials")
    ws.sheet_properties.tabColor = "4472C4"

    # Title
    ws.merge_cells("A1:E1")
    ws["A1"].value = "Historical & Projected Financials"
    ws["A1"].font = TITLE_FONT

    ws["A2"].value = "($ in millions)"
    ws["A2"].font = SUBTITLE_FONT

    header_row = 4
    data_start = 5
    cols = full_df.columns

    # Header row -- dark background, white text
    ws.cell(row=header_row, column=1, value="").font = HEADER_FONT
    ws.cell(row=header_row, column=1).fill = DARK_HEADER
    ws.cell(row=header_row, column=1).border = HEADER_BORDER

    for col_idx, col in enumerate(cols, start=2):
        cell = ws.cell(row=header_row, column=col_idx)
        cell.value = str(col.year) if hasattr(col, "year") else str(col)
        cell.font = HEADER_FONT
        cell.fill = DARK_HEADER
        cell.alignment = CENTER
        cell.border = HEADER_BORDER

    # Data rows
    for row_idx, metric in enumerate(full_df.index, start=data_start):
        is_pct = metric in PCT_ROWS
        is_bold = metric in BOLD_ROWS

        # Label cell
        label_cell = ws.cell(row=row_idx, column=1, value=metric)
        label_cell.font = BOLD_FONT if is_bold else LABEL_FONT
        label_cell.alignment = LEFT
        if is_bold:
            label_cell.fill = SUBTOTAL_FILL
            label_cell.border = Border(top=THIN, bottom=THIN)
        else:
            label_cell.border = CELL_BORDER

        for col_idx, col in enumerate(cols, start=2):
            raw = full_df.loc[metric, col]
            projected = not hasattr(col, "year")

            if is_pct:
                val = raw if not pd.isna(raw) else None
                fmt = "0.0%"
            else:
                val = _millions(raw)
                fmt = "#,##0"

            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.number_format = fmt
            cell.alignment = RIGHT
            cell.font = BOLD_FONT if is_bold else DATA_FONT

            if is_bold:
                cell.fill = SUBTOTAL_FILL if not projected else LIGHT_BLUE
                cell.border = Border(top=THIN, bottom=THIN)
            else:
                cell.border = CELL_BORDER
                if projected:
                    cell.fill = LIGHT_BLUE

    ws.column_dimensions["A"].width = 30
    for col_idx in range(2, len(cols) + 2):
        col_letter = ws.cell(row=header_row, column=col_idx).column_letter
        ws.column_dimensions[col_letter].width = 14


def _write_earnings(wb, filing_analysis):
    ws = wb.create_sheet("Earnings Analysis")
    ws.sheet_properties.tabColor = "548235"

    ws.merge_cells("A1:E1")
    ws["A1"].value = "SEC Filing Analysis"
    ws["A1"].font = TITLE_FONT

    # Filing metadata
    row = 3
    ws.cell(row=row, column=1, value="Filing Details").font = SECTION_FONT
    ws.cell(row=row, column=1).border = SECTION_BOTTOM
    ws.cell(row=row, column=2).border = SECTION_BOTTOM
    row += 1
    meta = [
        ("Filing Type", filing_analysis["filing_type"]),
        ("Filing Date", filing_analysis["filing_date"]),
        ("SEC URL", filing_analysis["url"]),
    ]
    for label, val in meta:
        ws.cell(row=row, column=1, value=label).font = LABEL_FONT
        ws.cell(row=row, column=2, value=val).font = DATA_FONT
        row += 1

    # Management Tone
    row += 1
    ws.cell(row=row, column=1, value="Management Tone").font = SECTION_FONT
    ws.cell(row=row, column=1).border = SECTION_BOTTOM
    ws.cell(row=row, column=2).border = SECTION_BOTTOM
    ws.cell(row=row, column=3).border = SECTION_BOTTOM
    row += 1

    tone = filing_analysis["tone_score"]
    pos_total = sum(filing_analysis["positive_counts"].values())
    neg_total = sum(filing_analysis["negative_counts"].values())
    tone_label = "N/A"
    if tone is not None:
        tone_label = "Net Positive" if tone > 0.55 else ("Net Negative" if tone < 0.45 else "Neutral")

    tone_items = [
        ("Tone Score", tone if tone is not None else "N/A", tone_label),
        ("Positive Keyword Mentions", pos_total, ""),
        ("Negative Keyword Mentions", neg_total, ""),
    ]
    for label, val, extra in tone_items:
        ws.cell(row=row, column=1, value=label).font = LABEL_FONT
        c = ws.cell(row=row, column=2, value=val)
        c.font = DATA_FONT
        if isinstance(val, float):
            c.number_format = "0.000"
        if extra:
            ec = ws.cell(row=row, column=3, value=extra)
            ec.font = Font(name="Calibri", size=10, bold=True,
                           color="006100" if "Positive" in extra else
                           ("9C0006" if "Negative" in extra else "808080"))
        row += 1

    # Keyword tables side by side
    row += 1
    ws.cell(row=row, column=1, value="Top Positive Keywords").font = SECTION_FONT
    ws.cell(row=row, column=1).fill = GREEN_FILL
    ws.cell(row=row, column=2).fill = GREEN_FILL
    ws.cell(row=row, column=4, value="Top Negative Keywords").font = SECTION_FONT
    ws.cell(row=row, column=4).fill = RED_FILL
    ws.cell(row=row, column=5).fill = RED_FILL
    row += 1

    sorted_pos = sorted(filing_analysis["positive_counts"].items(), key=lambda x: x[1], reverse=True)[:5]
    sorted_neg = sorted(filing_analysis["negative_counts"].items(), key=lambda x: x[1], reverse=True)[:5]

    kw_start = row
    for i, (word, count) in enumerate(sorted_pos):
        ws.cell(row=kw_start + i, column=1, value=word).font = DATA_FONT
        ws.cell(row=kw_start + i, column=2, value=count).font = DATA_FONT
        ws.cell(row=kw_start + i, column=2).alignment = RIGHT
    for i, (word, count) in enumerate(sorted_neg):
        ws.cell(row=kw_start + i, column=4, value=word).font = DATA_FONT
        ws.cell(row=kw_start + i, column=5, value=count).font = DATA_FONT
        ws.cell(row=kw_start + i, column=5).alignment = RIGHT

    # Guidance excerpts
    row = kw_start + max(len(sorted_pos), len(sorted_neg), 1) + 2
    ws.cell(row=row, column=1, value="Management Guidance Excerpts").font = SECTION_FONT
    ws.cell(row=row, column=1).border = SECTION_BOTTOM
    row += 1
    if filing_analysis["guidance_sentences"]:
        for i, sentence in enumerate(filing_analysis["guidance_sentences"], 1):
            ws.cell(row=row, column=1, value=f"{i}.").font = BOLD_FONT
            ws.cell(row=row, column=2, value=sentence).font = DATA_FONT
            row += 1
    else:
        ws.cell(row=row, column=1, value="No guidance sentences found in this filing.").font = SUBTITLE_FONT

    _set_col_widths(ws, {"A": 28, "B": 14, "C": 4, "D": 28, "E": 14})


def _write_dcf(wb, full_df, wacc, assumptions, dcf_results, current_price):
    ws = wb.create_sheet("DCF")
    ws.sheet_properties.tabColor = "BF8F00"

    proj_cols = [c for c in full_df.columns if not hasattr(c, "year")]
    projection_years = len(proj_cols)
    terminal_growth = assumptions["terminal_growth_rate"]
    total_debt_val = full_df.loc["Total Debt"].dropna().iloc[0]
    cash_val = full_df.loc["Cash And Cash Equivalents"].dropna().iloc[0]
    shares_val = dcf_results["equity_value"] / dcf_results["implied_price"]

    ws.merge_cells("A1:E1")
    ws["A1"].value = "DCF Valuation"
    ws["A1"].font = TITLE_FONT
    ws["A2"].value = "($ in millions, except per share data)"
    ws["A2"].font = SUBTITLE_FONT

    # --- FCF Bridge ---
    row = 4
    ws.cell(row=row, column=1, value="Free Cash Flow Build").font = SECTION_FONT
    ws.cell(row=row, column=1).border = SECTION_BOTTOM
    for ci in range(projection_years):
        ws.cell(row=row, column=ci + 2).border = SECTION_BOTTOM
    row += 1

    # Column headers
    ws.cell(row=row, column=1, value="").fill = DARK_HEADER
    ws.cell(row=row, column=1).border = HEADER_BORDER
    for ci, col in enumerate(proj_cols):
        cell = ws.cell(row=row, column=ci + 2, value=str(col))
        cell.font = HEADER_FONT
        cell.fill = DARK_HEADER
        cell.alignment = CENTER
        cell.border = HEADER_BORDER
    row += 1

    bridge_rows = [
        ("EBIT", "EBIT"),
        ("Tax Provision", "Less: Taxes"),
        ("NOPAT", "NOPAT"),
        ("Depreciation And Amortization", "+ D&A"),
        ("Capital Expenditure", "- CapEx"),
        ("Change In Working Capital", "- Change in NWC"),
        ("UFCF", "Unlevered Free Cash Flow"),
    ]
    for metric, label in bridge_rows:
        is_total = metric in ("NOPAT", "UFCF")
        lbl = ws.cell(row=row, column=1, value=label)
        lbl.font = BOLD_FONT if is_total else LABEL_FONT
        if is_total:
            lbl.fill = SUBTOTAL_FILL
            lbl.border = Border(top=THIN, bottom=THIN)
        else:
            lbl.border = CELL_BORDER

        for ci, col in enumerate(proj_cols):
            val = _millions(full_df.loc[metric, col])
            cell = ws.cell(row=row, column=ci + 2, value=val)
            cell.number_format = "#,##0"
            cell.alignment = RIGHT
            cell.font = BOLD_FONT if is_total else DATA_FONT
            if is_total:
                cell.fill = SUBTOTAL_FILL
                cell.border = Border(top=THIN, bottom=THIN)
            else:
                cell.border = CELL_BORDER
        row += 1

    # --- Discounting ---
    row += 1
    ws.cell(row=row, column=1, value="Discounted Cash Flows").font = SECTION_FONT
    ws.cell(row=row, column=1).border = SECTION_BOTTOM
    for ci in range(projection_years):
        ws.cell(row=row, column=ci + 2).border = SECTION_BOTTOM
    row += 1

    ws.cell(row=row, column=1, value="").fill = DARK_HEADER
    ws.cell(row=row, column=1).border = HEADER_BORDER
    for ci, col in enumerate(proj_cols):
        cell = ws.cell(row=row, column=ci + 2, value=str(col))
        cell.font = HEADER_FONT
        cell.fill = DARK_HEADER
        cell.alignment = CENTER
        cell.border = HEADER_BORDER
    row += 1

    disc_data = [
        ("UFCF", [_millions(full_df.loc["UFCF", c]) for c in proj_cols], "#,##0", False),
        ("Discount Factor", [1 / (1 + wacc) ** (i + 1) for i in range(projection_years)], "0.0000", False),
        ("PV of UFCF", [_millions(pv) for pv in dcf_results["pv_fcfs"]], "#,##0", True),
    ]
    for label, values, fmt, bold in disc_data:
        lbl = ws.cell(row=row, column=1, value=label)
        lbl.font = BOLD_FONT if bold else LABEL_FONT
        if bold:
            lbl.fill = SUBTOTAL_FILL
            lbl.border = Border(top=THIN, bottom=THIN)
        else:
            lbl.border = CELL_BORDER
        for ci, val in enumerate(values):
            cell = ws.cell(row=row, column=ci + 2, value=val)
            cell.number_format = fmt
            cell.alignment = RIGHT
            cell.font = BOLD_FONT if bold else DATA_FONT
            if bold:
                cell.fill = SUBTOTAL_FILL
                cell.border = Border(top=THIN, bottom=THIN)
            else:
                cell.border = CELL_BORDER
        row += 1

    # --- Terminal Value ---
    row += 1
    ws.cell(row=row, column=1, value="Terminal Value").font = SECTION_FONT
    ws.cell(row=row, column=1).border = SECTION_BOTTOM
    ws.cell(row=row, column=2).border = SECTION_BOTTOM
    row += 1

    tv_items = [
        ("Final Year UFCF", _millions(full_df.loc["UFCF", proj_cols[-1]]), "#,##0", False),
        ("Terminal Growth Rate", terminal_growth, "0.00%", False),
        ("WACC", wacc, "0.00%", False),
        ("Terminal Value", _millions(dcf_results["terminal_value"]), "#,##0", True),
        ("PV of Terminal Value", _millions(dcf_results["pv_terminal"]), "#,##0", True),
    ]
    for label, value, fmt, bold in tv_items:
        lbl = ws.cell(row=row, column=1, value=label)
        lbl.font = BOLD_FONT if bold else LABEL_FONT
        lbl.border = CELL_BORDER
        c = ws.cell(row=row, column=2, value=value)
        c.number_format = fmt
        c.font = BOLD_FONT if bold else DATA_FONT
        c.alignment = RIGHT
        c.border = CELL_BORDER
        row += 1

    # --- Equity Bridge ---
    row += 1
    ws.cell(row=row, column=1, value="Equity Bridge").font = SECTION_FONT
    ws.cell(row=row, column=1).border = SECTION_BOTTOM
    ws.cell(row=row, column=2).border = SECTION_BOTTOM
    row += 1

    eb_items = [
        ("Sum of PV of FCFs", _millions(sum(dcf_results["pv_fcfs"])), "#,##0", False),
        ("+ PV of Terminal Value", _millions(dcf_results["pv_terminal"]), "#,##0", False),
        ("= Enterprise Value", _millions(dcf_results["enterprise_value"]), "#,##0", True),
        ("- Total Debt", _millions(total_debt_val), "#,##0", False),
        ("+ Cash & Equivalents", _millions(cash_val), "#,##0", False),
        ("= Equity Value", _millions(dcf_results["equity_value"]), "#,##0", True),
        ("Shares Outstanding (M)", shares_val / SCALE, "#,##0.0", False),
        ("Implied Share Price", dcf_results["implied_price"], "$#,##0.00", True),
        ("Current Market Price", current_price, "$#,##0.00", False),
        ("Upside / Downside", (dcf_results["implied_price"] / current_price) - 1, "0.0%", True),
    ]
    for label, value, fmt, bold in eb_items:
        lbl = ws.cell(row=row, column=1, value=label)
        lbl.font = BOLD_FONT if bold else LABEL_FONT
        c = ws.cell(row=row, column=2, value=value)
        c.number_format = fmt
        c.alignment = RIGHT

        if label.startswith("="):
            lbl.fill = SUBTOTAL_FILL
            lbl.border = Border(top=THIN, bottom=THIN)
            c.fill = SUBTOTAL_FILL
            c.border = Border(top=THIN, bottom=THIN)
            c.font = BOLD_FONT
            lbl.font = BOLD_FONT
        elif bold:
            c.font = BOLD_FONT
            lbl.font = BOLD_FONT
            lbl.border = CELL_BORDER
            c.border = CELL_BORDER
        else:
            c.font = DATA_FONT
            lbl.border = CELL_BORDER
            c.border = CELL_BORDER
        row += 1

    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 16
    for ci in range(2, projection_years + 2):
        ws.column_dimensions[get_column_letter(ci)].width = 14


def _write_sensitivity(wb, sens_df, sens2_df, current_price):
    ws = wb.create_sheet("Sensitivity")
    ws.sheet_properties.tabColor = "C00000"

    ws.merge_cells("A1:F1")
    ws["A1"].value = "Sensitivity Analysis"
    ws["A1"].font = TITLE_FONT

    def write_table(ws, df, title, corner_label, start_row):
        ws.cell(row=start_row, column=1, value=title).font = SECTION_FONT
        ws.cell(row=start_row, column=1).border = SECTION_BOTTOM

        header_row = start_row + 1
        # Corner cell
        cell = ws.cell(row=header_row, column=1, value=corner_label)
        cell.font = HEADER_FONT
        cell.fill = DARK_HEADER
        cell.alignment = CENTER
        cell.border = HEADER_BORDER

        # Column headers
        for ci, col_name in enumerate(df.columns):
            cell = ws.cell(row=header_row, column=ci + 2, value=col_name)
            cell.font = HEADER_FONT
            cell.fill = DARK_HEADER
            cell.alignment = CENTER
            cell.border = HEADER_BORDER

        # Data
        for ri, idx in enumerate(df.index):
            # Row header
            rh = ws.cell(row=header_row + 1 + ri, column=1, value=idx)
            rh.font = BOLD_FONT
            rh.fill = SECTION_FILL
            rh.alignment = CENTER
            rh.border = CELL_BORDER

            for ci, col_name in enumerate(df.columns):
                val = df.loc[idx, col_name]
                cell = ws.cell(row=header_row + 1 + ri, column=ci + 2, value=val)
                cell.number_format = "$#,##0.00"
                cell.alignment = RIGHT
                cell.font = DATA_FONT
                cell.border = CELL_BORDER

                if val >= current_price:
                    cell.fill = GREEN_FILL
                elif val < current_price * 0.9:
                    cell.fill = RED_FILL
                else:
                    cell.fill = WHITE_FILL

        return header_row + 1 + len(df.index) + 2

    next_row = write_table(ws, sens_df, "WACC vs Terminal Growth Rate", "WACC \\ TGR", 3)
    write_table(ws, sens2_df, "Revenue Growth vs Gross Margin", "Growth \\ GM", next_row)

    ws.column_dimensions["A"].width = 14
    for ci in range(2, max(len(sens_df.columns), len(sens2_df.columns)) + 2):
        ws.column_dimensions[get_column_letter(ci)].width = 14


def generate_excel(
    ticker, company_name, current_price, implied_price, wacc,
    assumptions, full_df, dcf_results, filing_analysis,
    sens_df, sens2_df, filename=None
):
    wb = Workbook()

    _write_summary(wb, ticker, company_name, current_price, implied_price, wacc, assumptions)
    _write_financials(wb, full_df)
    _write_earnings(wb, filing_analysis)
    _write_dcf(wb, full_df, wacc, assumptions, dcf_results, current_price)
    _write_sensitivity(wb, sens_df, sens2_df, current_price)

    if filename is None:
        filename = f"{ticker}_dcf_model.xlsx"

    wb.save(filename)
    print(f"Saved to {filename}")


def generate_excel_bytes(
    ticker, company_name, current_price, implied_price, wacc,
    assumptions, full_df, dcf_results, filing_analysis,
    sens_df, sens2_df,
):
    """Return the workbook as in-memory bytes for download."""
    wb = Workbook()
    _write_summary(wb, ticker, company_name, current_price, implied_price, wacc, assumptions)
    _write_financials(wb, full_df)
    _write_earnings(wb, filing_analysis)
    _write_dcf(wb, full_df, wacc, assumptions, dcf_results, current_price)
    _write_sensitivity(wb, sens_df, sens2_df, current_price)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()
